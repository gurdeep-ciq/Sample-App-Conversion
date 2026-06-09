"""
profiling.py — file parsing, sheet detection, column profiling, and the
deterministic pipeline executor.

This is the Python port of the original in-browser heuristic engine, extended to:
  * profile ALL sheets in a workbook (not just the largest one), and
  * expose a shared `apply_pipeline` executor used by the Dagster assets.

No LLM calls. Pure heuristics + stdlib + openpyxl.
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import re
from typing import Any, Optional

try:
    import openpyxl  # type: ignore
except Exception:  # pragma: no cover
    openpyxl = None


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def snake(s: Any) -> str:
    s = "" if s is None else str(s)
    s = re.sub(r"[^a-z0-9]+", "_", s.strip().lower()).strip("_")
    s = re.sub(r"^(\d)", r"col_\1", s)
    return s or "column"


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


# --------------------------------------------------------------------------
# workbook reading -> list[{name, aoa}]
# --------------------------------------------------------------------------
def read_workbook(path: str, filename: Optional[str] = None) -> list[dict]:
    name = (filename or path).lower()
    if name.endswith((".csv", ".tsv", ".txt")):
        return [_read_delimited(path)]
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to read .xlsx files")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        aoa = [list(row) for row in ws.iter_rows(values_only=True)]
        # trim fully-empty trailing rows
        while aoa and all(_is_blank(c) for c in aoa[-1]):
            aoa.pop()
        sheets.append({"name": ws.title, "aoa": aoa})
    wb.close()
    return sheets


def _read_delimited(path: str) -> dict:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        text = f.read()
    first = text.split("\n", 1)[0]
    delim = "\t" if first.count("\t") > first.count(",") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    aoa = []
    for r in rows:
        out = []
        for v in r:
            if v == "":
                out.append(None)
            elif re.fullmatch(r"-?[\d.]+", v) and v.count(".") <= 1:
                try:
                    out.append(float(v) if "." in v else int(v))
                except ValueError:
                    out.append(v)
            else:
                out.append(v)
        aoa.append(out)
    return {"name": "Sheet1", "aoa": aoa}


# --------------------------------------------------------------------------
# sheet classification + header detection
# --------------------------------------------------------------------------
_META_NAME = re.compile(
    r"^(filters?|extra ?info|.*info|cover|notes?|readme|summary|legend|parameters?|settings?)$",
    re.I,
)


def detect_header_row(aoa: list[list]) -> int:
    best_row, best_score = 0, float("-inf")
    for r in range(min(len(aoa), 25)):
        row = aoa[r] or []
        non_empty = [c for c in row if not _is_blank(c)]
        if len(non_empty) < 2:
            continue
        str_cells = [
            c for c in non_empty
            if isinstance(c, str) and 0 < len(c) < 60
            and not re.fullmatch(r"-?\d+(\.\d+)?", c.strip())
        ]
        below = [rr for rr in aoa[r + 1 : r + 6] if rr and any(not _is_blank(c) for c in rr)]
        if not below:
            continue
        below_fill = sum(len([c for c in rr if not _is_blank(c)]) for rr in below) / len(below)
        if below_fill < 2:
            continue
        score = len(str_cells) * 2 + min(below_fill, len(non_empty)) - r * 0.6
        if score > best_score:
            best_row, best_score = r, score
    return best_row


def classify_sheet(name: str, aoa: list[list]) -> dict:
    aoa = [list(r) if r else [] for r in (aoa or [])]
    non_empty_rows = [r for r in aoa if any(not _is_blank(c) for c in r)]
    header_row = detect_header_row(aoa) if aoa else 0
    width = len([c for c in (aoa[header_row] if header_row < len(aoa) else []) if not _is_blank(c)])
    data_rows = max(0, len(aoa) - header_row - 1)

    looks_kv = width <= 2 and len(non_empty_rows) >= 2
    too_small = width < 3 or data_rows < 1
    meta_named = bool(_META_NAME.match(name.strip())) and width < 4
    kind = "metadata" if (looks_kv or meta_named or too_small) else "data"
    return {
        "name": name, "aoa": aoa, "kind": kind, "headerRow": header_row,
        "width": width, "dataRows": data_rows,
    }


# --------------------------------------------------------------------------
# column type inference
# --------------------------------------------------------------------------
_RE_CURRENCY = re.compile(r"^-?\$\s?[\d,]+(\.\d+)?$")
_RE_CURRENCY2 = re.compile(r"^-?[\d,]+\.\d{2}$")
_RE_NUM = re.compile(r"^-?[\d,]+(\.\d+)?$")
_RE_DATE = re.compile(r"^\d{4}-\d{1,2}-\d{1,2}|^\d{1,2}/\d{1,2}/\d{2,4}")
_RE_TIME = re.compile(r"^\d{1,2}:\d{2}(:\d{2})?(\s?[ap]m)?$", re.I)
_RE_LEADZERO = re.compile(r"^0\d+$")


def _frac(values, pred) -> float:
    return (sum(1 for v in values if pred(v)) / len(values)) if values else 0.0


def infer_column(values: list) -> dict:
    vals = [v for v in values if not _is_blank(v)]
    total = len(values)
    null_rate = (total - len(vals)) / total if total else 1.0
    distinct = len({str(v) for v in vals})
    if not vals:
        return dict(type="empty", nullRate=1.0, distinct=0, sample=[], allZero=False,
                    leadingZero=False, bigIntNum=False, isCurrency=False, constant=True, numeric=False)

    def is_num(v):
        return isinstance(v, (int, float)) and not isinstance(v, bool) or (
            isinstance(v, str) and bool(_RE_NUM.match(v.strip())))

    def is_date(v):
        return isinstance(v, (_dt.date, _dt.datetime)) or (isinstance(v, str) and bool(_RE_DATE.match(v)))

    def is_time(v):
        return isinstance(v, _dt.time) or (isinstance(v, str) and bool(_RE_TIME.match(v.strip())))

    def is_curr(v):
        s = str(v).strip()
        return bool(_RE_CURRENCY.match(s) or _RE_CURRENCY2.match(s)) and not re.fullmatch(r"-?\d+", s)

    isdate = _frac(vals, is_date) > 0.8
    istime = (not isdate) and _frac(vals, is_time) > 0.8
    iscurr = _frac(vals, is_curr) > 0.6 or _frac(vals, lambda v: isinstance(v, str) and "$" in v) > 0.5
    numeric = _frac(vals, is_num) > 0.85

    def _num(v):
        try:
            return float(re.sub(r"[,$%\s]", "", str(v)))
        except ValueError:
            return None

    all_zero = numeric and all((_num(v) == 0) for v in vals)
    leading_zero = _frac(vals, lambda v: bool(_RE_LEADZERO.match(str(v).strip()))) > 0.5
    big_int_num = numeric and all(
        isinstance(v, (int, float)) and not isinstance(v, bool)
        and float(v).is_integer() and abs(v) >= 1e8 for v in vals)
    constant = distinct == 1

    if isdate:
        typ = "date"
    elif istime:
        typ = "time"
    elif iscurr:
        typ = "currency"
    elif numeric:
        typ = "number"
    else:
        typ = "string"

    return dict(type=typ, nullRate=round(null_rate, 4), distinct=distinct,
                sample=[_jsonable(v) for v in vals[:5]], allZero=all_zero, leadingZero=leading_zero,
                bigIntNum=big_int_num, constant=constant, isCurrency=iscurr, numeric=numeric)


def _jsonable(v):
    if isinstance(v, (_dt.date, _dt.datetime, _dt.time)):
        return v.isoformat()
    return v


_MEASURE_HINT = re.compile(r"(qty|quantity|units?|sold|cases?|bottles?|amount|amt|revenue|sales|net|gross|price|cost|fob|msrp|total|count|sum|volume|liters?)", re.I)
_ID_HINT = re.compile(r"(sku|upc|ean|gtin|barcode|code|id$|_id|key|account|product|customer)", re.I)
_DATE_HINT = re.compile(r"(date|day|month|year|period)", re.I)
# boundary-safe: leading start/non-letter, trailing non-letter/end (so "Employee_Matt" matches, "Sampler" does not)
INTERNAL_RE = re.compile(r"(^|[^a-z])(employee|salesrep|sales ?rep|sample|comp|complimentary|donation|staff|internal|supplier|house ?account|admin)([^a-z]|$)", re.I)


def role_for(name: str, col: dict) -> str:
    if col["type"] == "empty" or col["constant"] or col["allZero"]:
        return "ignore"
    if col["type"] in ("date", "time") or _DATE_HINT.search(name) or re.search(r"\btime\b", name, re.I):
        return "timestamp"
    if _ID_HINT.search(name) or col["leadingZero"] or col["bigIntNum"]:
        return "id"
    if col["type"] in ("number", "currency"):
        return "measure"
    return "dimension"


def target_type(name: str, col: dict, role: str) -> str:
    if role == "id":
        return "string"
    if col["type"] == "currency":
        return "decimal"
    if role == "timestamp" or col["type"] == "date":
        return "date"
    if col["type"] == "time":
        return "time"
    if col["type"] == "number":
        s0 = col["sample"][0] if col["sample"] else 0
        try:
            is_int = float(s0).is_integer()
        except (ValueError, TypeError):
            is_int = False
        return "integer" if (is_int and not _MEASURE_HINT.search(name)) else "decimal"
    return "string"


# --------------------------------------------------------------------------
# build a profile for a single data sheet
# --------------------------------------------------------------------------
def build_profile(sheet: dict) -> dict:
    aoa, header_row = sheet["aoa"], sheet["headerRow"]
    raw_headers = [
        (f"column_{i+1}" if _is_blank(h) else str(h).strip())
        for i, h in enumerate(aoa[header_row] if header_row < len(aoa) else [])
    ]
    data_rows = [r for r in aoa[header_row + 1:] if r and any(not _is_blank(c) for c in r)]
    columns = []
    for i, h in enumerate(raw_headers):
        col_vals = [r[i] if i < len(r) else None for r in data_rows]
        info = infer_column(col_vals)
        role = role_for(h, info)
        columns.append({
            "index": i, "source": h, "canonical": snake(h),
            **info, "role": role, "ttype": target_type(h, info, role),
        })
    return {
        "sheetName": sheet["name"], "headerRow": header_row, "columns": columns,
        "rowCount": len(data_rows), "rawHeaders": raw_headers,
        "_dataRows": data_rows,  # kept in-memory for rule impact / execution; stripped before JSON
    }


# --------------------------------------------------------------------------
# rule + schema generation (per sheet)
# --------------------------------------------------------------------------
def generate_rules_and_schema(sheet: dict, profile: dict) -> dict:
    rules, rid = [], [0]

    def add(**r):
        rid[0] += 1
        rules.append({"id": f"r{rid[0]}", "enabled": r.pop("enabled", True), **r})

    cols = profile["columns"]
    data_rows = profile["_dataRows"]

    add(kind="select_sheet", title=f'Use sheet "{profile["sheetName"]}" as a data source',
        confidence=0.95, spec={"sheet": profile["sheetName"]},
        why="Tabular data sheet.")
    if profile["headerRow"] > 0:
        add(kind="header_row", title=f'Treat row {profile["headerRow"]+1} as the header',
            confidence=0.9, spec={"headerRow": profile["headerRow"]},
            why="Rows above are titles/metadata.")

    for c in cols:
        if c["role"] == "id" and (c["bigIntNum"] or c["numeric"]) and not c["leadingZero"]:
            add(kind="cast", title=f'Store "{c["source"]}" as text to preserve the full code',
                confidence=0.9, spec={"column": c["source"], "to": "string"},
                why="Identifier (SKU/UPC/code) — numeric storage risks precision loss.")
        if c["isCurrency"]:
            add(kind="cast", title=f'Parse "{c["source"]}" as a number (strip $ and commas)',
                confidence=0.9, spec={"column": c["source"], "to": "currency"},
                why=f'Sample: {", ".join(str(s) for s in c["sample"][:2])}')
        if c["type"] == "date":
            add(kind="cast", title=f'Normalize "{c["source"]}" to ISO date', confidence=0.85,
                spec={"column": c["source"], "to": "date"}, why="Mixed date formats normalize.")

    date_col = next((c for c in cols if re.fullmatch(r"date", c["source"], re.I)
                     or (c["type"] == "date" and re.search(r"date", c["source"], re.I))), None)
    time_col = next((c for c in cols if re.fullmatch(r"time", c["source"], re.I) or c["type"] == "time"), None)
    if date_col and time_col and date_col["index"] != time_col["index"]:
        add(kind="combine_datetime",
            title=f'Combine "{date_col["source"]}" + "{time_col["source"]}" into one timestamp',
            confidence=0.8, spec={"dateCol": date_col["source"], "timeCol": time_col["source"],
                                  "target": "event_timestamp"}, why="Date and time are split.")

    for c in cols:
        if c["role"] == "ignore":
            reason = "always empty" if c["type"] == "empty" else "always zero" if c["allZero"] else "single constant value"
            add(kind="drop_column", title=f'Drop "{c["source"]}" ({reason})', confidence=0.85,
                spec={"column": c["source"]}, why=f"Carries no information ({reason}).")

    internal_col = None
    for c in cols:
        if c["type"] != "string":
            continue
        probe = [c["sample"]] + [[r[c["index"]] if c["index"] < len(r) else None for r in data_rows[:300]]]
        if any(v and INTERNAL_RE.search(str(v)) for grp in probe for v in grp):
            internal_col = c
            break
    if internal_col:
        toks = sorted({INTERNAL_RE.search(str(r[internal_col["index"]])).group(2).strip()
                       for r in data_rows
                       if internal_col["index"] < len(r) and r[internal_col["index"]]
                       and INTERNAL_RE.search(str(r[internal_col["index"]]))})[:6]
        add(kind="filter", title=f'Drop internal/non-customer rows in "{internal_col["source"]}"',
            confidence=0.7, spec={"column": internal_col["source"], "op": "regex_not",
                                  "value": INTERNAL_RE.pattern, "label": ", ".join(toks)},
            why=f"Found values like: {', '.join(toks)} — typically employee/sample/internal rows.")

    qty_col = next((c for c in cols if c["type"] == "number"
                    and re.search(r"(qty|quantity|units?|sold|cases?)", c["source"], re.I)), None)
    if qty_col:
        add(kind="filter", title=f'Drop rows where "{qty_col["source"]}" is zero or blank',
            confidence=0.6, enabled=False, spec={"column": qty_col["source"], "op": "gt", "value": 0},
            why="Zero-quantity rows are often placeholders. Off by default.")

    name_col = next((c for c in cols if re.search(r"(customer|account|name|store|outlet|client)", c["source"], re.I)
                     and c["type"] == "string" and c["role"] != "id"), None)
    if name_col and name_col["nullRate"] > 0.001:
        add(kind="filter", title=f'Drop rows with a blank "{name_col["source"]}"', confidence=0.7,
            spec={"column": name_col["source"], "op": "not_empty"},
            why=f'{round(name_col["nullRate"]*100)}% of rows have no value here.')

    add(kind="dedupe", title="Drop exact duplicate rows", confidence=0.6, enabled=False, spec={},
        why="Off by default.")

    # schema
    dropped = {r["spec"]["column"] for r in rules if r["kind"] == "drop_column" and r["enabled"]}
    schema_cols = []
    for c in cols:
        if c["source"] in dropped or c["role"] == "ignore":
            continue
        schema_cols.append({"name": c["canonical"], "type": c["ttype"], "role": c["role"],
                            "source": c["source"], "include": True})
    if date_col and time_col:
        drop = {date_col["source"], time_col["source"]}
        schema_cols = [c for c in schema_cols if c["source"] not in drop]
        schema_cols.insert(0, {"name": "event_timestamp", "type": "timestamp", "role": "timestamp",
                               "source": f'{date_col["source"]}+{time_col["source"]}', "include": True})

    table = snake(profile["sheetName"])
    return {"rules": rules, "schema": {"table": table, "columns": schema_cols},
            "formatGuess": _guess_format(profile)}


def _guess_format(profile: dict) -> str:
    headers = " ".join(profile["rawHeaders"]).lower()
    if re.search(r"supplier key|supplier's sku|product key", headers):
        return "vendor-multi-sheet-report"
    if re.search(r"time zone|price point|modifiers applied", headers):
        return "pos-transaction-export"
    if re.search(r"importer|depletion|account region|sales rep", headers):
        return "distributor-depletion"
    return "flat-tabular"


# --------------------------------------------------------------------------
# pipeline executor (shared by Dagster assets)
# --------------------------------------------------------------------------
def _to_number(v):
    if _is_blank(v):
        return None
    try:
        return float(re.sub(r"[,$%\s]", "", str(v)))
    except ValueError:
        return None


def _to_iso_date(v):
    if isinstance(v, (_dt.date, _dt.datetime)):
        return v.date().isoformat() if isinstance(v, _dt.datetime) else v.isoformat()
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return f"{m[1]}-{int(m[2]):02d}-{int(m[3]):02d}"
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})", s)
    if m:
        y = m[3] if len(m[3]) == 4 else "20" + m[3]
        return f"{y}-{int(m[1]):02d}-{int(m[2]):02d}"
    return s


def _to_time_of_day(v):
    """Normalize a time cell to HH:MM:SS. Spreadsheet readers hand back times as
    datetime.time, as a datetime on the 1899-12-31 epoch, as an Excel serial
    fraction (fraction of a 24h day), or as a string — handle all of them."""
    if _is_blank(v):
        return "00:00:00"
    if isinstance(v, (_dt.datetime, _dt.time)):
        return v.strftime("%H:%M:%S")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        s = round((float(v) - int(v)) * 86400)
        return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"
    m = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", str(v))
    return f"{int(m.group(1)):02d}:{m.group(2)}:{m.group(3) or '00'}" if m else "00:00:00"


def apply_pipeline(profile: dict, rules: list[dict], schema_cols: list[dict]) -> dict:
    """Run the enabled rules against the sheet's data rows. Returns cleaned rows,
    dropped rows (with reason), and per-column DQ stats."""
    enabled = [r for r in rules if r.get("enabled")]
    headers = profile["rawHeaders"]
    rows = []
    for r in profile["_dataRows"]:
        rows.append({h: (r[i] if i < len(r) else None) for i, h in enumerate(headers)})
    input_count = len(rows)
    dropped = []

    # combine runs BEFORE casts so it reads the original Date/Time cells, not
    # cast-mutated values, and uses robust time-of-day parsing.
    for r in [x for x in enabled if x["kind"] == "combine_datetime"]:
        sp = r["spec"]
        for o in rows:
            d = _to_iso_date(o.get(sp["dateCol"]))
            t = _to_time_of_day(o.get(sp["timeCol"]))
            o[sp["target"]] = f"{d}T{t}" if d else None

    for r in [x for x in enabled if x["kind"] == "cast"]:
        col, to = r["spec"]["column"], r["spec"]["to"]
        for o in rows:
            if _is_blank(o.get(col)):
                continue
            if to == "string":
                v = o[col]
                o[col] = (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
            elif to == "currency":
                o[col] = _to_number(o[col])
            elif to == "date":
                o[col] = _to_iso_date(o[col])

    filters = [x for x in enabled if x["kind"] == "filter"]
    kept = []
    for o in rows:
        drop_reason = None
        for r in filters:
            sp = r["spec"]
            v = o.get(sp["column"])
            op = sp["op"]
            keep = True
            if op == "not_empty":
                keep = not _is_blank(v)
            elif op == "gt":
                keep = _to_number(v) is not None and _to_number(v) > sp["value"]
            elif op == "regex_not":
                keep = not (v is not None and re.search(sp["value"], str(v), re.I))
            elif op == "not_in":
                keep = str(v).lower() not in [str(x).lower() for x in sp.get("value", [])]
            if not keep:
                drop_reason = r["title"]
                break
        (kept if drop_reason is None else None)
        if drop_reason is None:
            kept.append(o)
        else:
            dropped.append({"row": o, "rule": drop_reason})
    rows = kept

    if any(x["kind"] == "dedupe" for x in enabled):
        seen, out = set(), []
        for o in rows:
            k = repr(sorted(o.items(), key=lambda kv: kv[0]))
            if k in seen:
                dropped.append({"row": o, "rule": "Exact duplicate"})
            else:
                seen.add(k)
                out.append(o)
        rows = out

    # project onto schema
    out_rows = []
    for o in rows:
        rec = {}
        for c in [c for c in schema_cols if c["include"]]:
            src = c.get("source", "")
            if "+" in src:
                rec[c["name"]] = o.get(c["name"]) or o.get("event_timestamp")
            else:
                rec[c["name"]] = _jsonable(o.get(src))
        out_rows.append(rec)

    dq = []
    for c in [c for c in schema_cols if c["include"]]:
        vals = [r.get(c["name"]) for r in out_rows]
        nulls = sum(1 for v in vals if _is_blank(v))
        dq.append({"name": c["name"], "type": c["type"],
                   "nullPct": round(nulls / len(out_rows) * 100, 1) if out_rows else 0.0,
                   "distinct": len({str(v) for v in vals})})

    return {"inputCount": input_count, "keptCount": len(out_rows),
            "droppedCount": len(dropped), "rows": out_rows,
            "dropped": [{"reason": d["rule"], "sample": _first_val(d["row"])} for d in dropped[:200]],
            "dq": dq}


def _first_val(row: dict):
    for v in row.values():
        if not _is_blank(v):
            return _jsonable(v)
    return None


# --------------------------------------------------------------------------
# top-level: profile an entire workbook (ALL sheets)
# --------------------------------------------------------------------------
def profile_workbook(path: str, filename: Optional[str] = None) -> dict:
    raw_sheets = read_workbook(path, filename)
    classified = [classify_sheet(s["name"], s["aoa"]) for s in raw_sheets]
    sheets_out = []
    for sh in classified:
        if sh["kind"] != "data":
            sheets_out.append({"name": sh["name"], "kind": "metadata", "headerRow": sh["headerRow"],
                               "rowCount": sh["dataRows"], "columns": [], "rules": [], "schema": None,
                               "formatGuess": None})
            continue
        prof = build_profile(sh)
        gen = generate_rules_and_schema(sh, prof)
        sheets_out.append({
            "name": prof["sheetName"], "kind": "data", "headerRow": prof["headerRow"],
            "rowCount": prof["rowCount"],
            "columns": [{k: v for k, v in c.items()} for c in prof["columns"]],
            "rules": gen["rules"], "schema": gen["schema"], "formatGuess": gen["formatGuess"],
            "_profile": prof,  # kept in-memory only
        })
    return {"fileName": filename or path, "sheets": sheets_out}
